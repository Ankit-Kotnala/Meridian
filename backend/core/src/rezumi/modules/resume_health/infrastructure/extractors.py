"""Bounded local PDF/DOCX validation and deterministic text extraction."""

import asyncio
import html
import re
import stat
from collections.abc import Iterable
from pathlib import Path, PurePosixPath
from uuid import UUID, uuid5
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from rezumi.modules.resume_health.application.models import (
    DocumentLimits,
    ExtractedBlock,
    ExtractionResult,
    SourceSpanView,
)
from rezumi.modules.resume_health.domain import ResumeMediaType
from rezumi.modules.resume_health.domain.errors import UnsafeDocument
from rezumi.modules.resume_health.infrastructure.hyperlinks import (
    append_hyperlink_targets,
    safe_http_hyperlink,
)
from rezumi.modules.resume_health.infrastructure.layout import analyze_local_layout
from rezumi.modules.resume_health.infrastructure.pdf_reading_order import extract_pdf_pages

PARSER_VERSION = "rezumi-local-parser/1.3.0"
_PDF_ACTIVE_CONTENT = (b"/JavaScript", b"/Launch", b"/EmbeddedFile")
_PDF_JS_NAME = re.compile(rb"/JS(?=[/\s\[\]<>()])")
_BLOCK_NAMESPACE = UUID("5f80ce6a-096a-44e9-b4d1-335f2da30c32")
_BIDI_CONTROLS = dict.fromkeys(
    [
        *range(0x202A, 0x202F),
        *range(0x2066, 0x206A),
        0x200E,
        0x200F,
        0x061C,
    ]
)
_BULLET_PREFIX = re.compile(r"^(?:[-*•▪◦·\u2022\uf0b7\uf0a7\u25aa\u25cf]|\d+[.)])\s+")
# PDF icon fonts and contact glyphs that often precede phone/email/link text.
_CONTACT_ICON_GLYPHS = re.compile(
    r"[\u2640\u2642\u260e\u2709\u2706\uf095\uf003\uf0e0\uf0ac\uf099\uf09a\uf099]+"
)
# Resume templates sometimes emit "/linkedin" before "linkedin.com/..." in one line.
# The same artifact is produced by code/professional profile icons. Social
# destinations are intentionally not normalized into declared career links.
_PROFILE_HOST_LABELS = (
    "linkedin",
    "github",
    "gitlab",
    "bitbucket",
)
_PROFILE_HOST_LABEL_PATTERN = "|".join(_PROFILE_HOST_LABELS)
_DUPLICATE_PROFILE_HOST = re.compile(
    rf"(?i)(?:/+(?:www\.)?|(?<![\w./])(?:www\.)?)"
    rf"(?P<label>{_PROFILE_HOST_LABEL_PATTERN})(?P=label)\.com"
)
_SLASH_BEFORE_PROFILE_HOST = re.compile(
    rf"(?<![:/])/+(?=(?:www\.)?(?:{_PROFILE_HOST_LABEL_PATTERN})\.com)",
    re.IGNORECASE,
)
_PHONE_LABEL_GLUE = re.compile(r"(?i)\bphone(?=\s*\+?\d)")
_DOCX_RELATIONSHIP_TAG = re.compile(rb"<Relationship\b[^>]*>", re.IGNORECASE)
_DOCX_XML_ATTRIBUTE = re.compile(rb"\b([A-Za-z][\w.-]*)\s*=\s*(['\"])(.*?)\2", re.DOTALL)


class LocalDocumentExtractor:
    async def extract(
        self, path: Path, media_type: str, limits: DocumentLimits
    ) -> ExtractionResult:
        try:
            return await asyncio.wait_for(
                asyncio.to_thread(self._extract, path, media_type, limits),
                limits.processing_timeout_seconds,
            )
        except TimeoutError as exc:
            raise UnsafeDocument("document_processing_timeout") from exc

    def _extract(self, path: Path, media_type: str, limits: DocumentLimits) -> ExtractionResult:
        return extract_local_document(path, media_type, limits)


class LocalDocumentTextExtractor:
    async def extract_text(
        self, path: Path, media_type: str, limits: DocumentLimits
    ) -> ExtractionResult:
        try:
            return await asyncio.wait_for(
                asyncio.to_thread(extract_local_text, path, media_type, limits),
                limits.processing_timeout_seconds,
            )
        except TimeoutError as exc:
            raise UnsafeDocument("document_processing_timeout") from exc


def extract_local_document(path: Path, media_type: str, limits: DocumentLimits) -> ExtractionResult:
    """Run local text and layout adapters inside the caller's isolation boundary."""
    return analyze_local_layout(extract_local_text(path, media_type, limits))


def extract_local_text(path: Path, media_type: str, limits: DocumentLimits) -> ExtractionResult:
    """Run deterministic local text extraction without layout policy."""
    if path.stat().st_size > limits.max_upload_bytes:
        raise UnsafeDocument("upload_too_large")
    if media_type == ResumeMediaType.PDF.value:
        return _extract_pdf(path, limits)
    if media_type == ResumeMediaType.DOCX.value:
        return _extract_docx(path, limits)
    raise UnsafeDocument("unsupported_document_type")


def _extract_pdf(path: Path, limits: DocumentLimits) -> ExtractionResult:
    raw = path.read_bytes()
    raw_head = raw[:1024]
    if not raw_head.startswith(b"%PDF-"):
        raise UnsafeDocument("document_signature_mismatch")
    if b"PK\x03\x04" in raw:
        raise UnsafeDocument("polyglot_document_rejected")
    if any(token in raw for token in _PDF_ACTIVE_CONTENT) or _PDF_JS_NAME.search(raw):
        raise UnsafeDocument("active_content_rejected")
    with path.open("rb") as source:
        source.seek(max(0, path.stat().st_size - 2048))
        if b"%%EOF" not in source.read():
            raise UnsafeDocument("malformed_pdf")
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            raise UnsafeDocument("encrypted_document")
        if not 1 <= len(reader.pages) <= limits.max_pdf_pages:
            raise UnsafeDocument("pdf_page_limit_exceeded")
        page_texts, position_signals = extract_pdf_pages(list(reader.pages))
        layout_signals: list[str] = list(position_signals)
        sanitized_pages: list[str] = []
        for extracted in page_texts:
            sanitized, bidi_removed = _sanitize_extracted_text(extracted.replace("\x00", ""))
            sanitized_pages.append(sanitized)
            if bidi_removed:
                layout_signals.append("bidirectional_controls_present")
        page_texts = sanitized_pages
    except UnsafeDocument:
        raise
    except Exception as exc:
        raise UnsafeDocument("malformed_pdf") from exc
    plain_text = "\n\f\n".join(page_texts).strip()
    if len(plain_text) > limits.max_extracted_characters:
        raise UnsafeDocument("extracted_text_limit_exceeded")
    blocks = _blocks_from_pages(page_texts, limits.max_extracted_blocks)
    image_only = len(plain_text.strip()) < 40
    warnings = ("image_only_pdf",) if image_only else ()
    return ExtractionResult(
        plain_text=plain_text,
        reading_order=blocks,
        page_count=len(page_texts),
        image_only=image_only,
        warnings=warnings,
        parser_version=PARSER_VERSION,
        layout_signals=tuple(dict.fromkeys(layout_signals)),
    )


def _extract_docx(path: Path, limits: DocumentLimits) -> ExtractionResult:
    raw = path.read_bytes()
    if not raw.startswith(b"PK\x03\x04"):
        raise UnsafeDocument("document_signature_mismatch")
    if b"%PDF-" in raw:
        raise UnsafeDocument("polyglot_document_rejected")
    try:
        with ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_archive_entries:
                raise UnsafeDocument("archive_entry_limit_exceeded")
            total_uncompressed = 0
            total_compressed = 0
            names: set[str] = set()
            for info in infos:
                normalized = PurePosixPath(info.filename.replace("\\", "/"))
                if normalized.is_absolute() or ".." in normalized.parts:
                    raise UnsafeDocument("archive_path_traversal")
                mode = info.external_attr >> 16
                if mode and stat.S_ISLNK(mode):
                    raise UnsafeDocument("archive_symlink_rejected")
                if info.flag_bits & 0x1:
                    raise UnsafeDocument("encrypted_document")
                total_uncompressed += info.file_size
                total_compressed += max(1, info.compress_size)
                names.add(info.filename.casefold())
            if total_uncompressed > limits.max_archive_uncompressed_bytes:
                raise UnsafeDocument("archive_expansion_limit_exceeded")
            if total_uncompressed > total_compressed * limits.max_archive_ratio:
                raise UnsafeDocument("archive_ratio_limit_exceeded")
            if "word/vbaproject.bin" in names or any("macroenabled" in name for name in names):
                raise UnsafeDocument("macro_document_rejected")
            if any(
                name.startswith("word/activex/") or name.startswith("word/embeddings/")
                for name in names
            ):
                raise UnsafeDocument("embedded_object_rejected")
            if "[content_types].xml" not in names or "word/document.xml" not in names:
                raise UnsafeDocument("invalid_docx_package")
            for info in infos:
                if info.filename.casefold().endswith(".rels"):
                    relationships = archive.read(info)
                    _validate_docx_relationships(relationships)
            if archive.testzip() is not None:
                raise UnsafeDocument("malformed_docx")
    except UnsafeDocument:
        raise
    except (BadZipFile, OSError) as exc:
        raise UnsafeDocument("malformed_docx") from exc

    try:
        document = Document(str(path))
        layout_signals: list[str] = []
        values: list[tuple[str, str]] = []
        # document.paragraphs and document.tables are separate views; reading
        # them in two loops moves every table after every paragraph. Resumes
        # routinely use tables for headers and experience rows, so keep the
        # package's body order and preserve Word list semantics that are not
        # represented by a visible bullet character in paragraph.text.
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                value, recovered = _docx_value_with_hyperlinks(item.text, item)
                sanitized, bidi_removed = _sanitize_extracted_text(value.replace("\x00", ""))
                values.append(
                    (
                        "bullet" if _is_docx_list_paragraph(item) else "text",
                        sanitized.strip(),
                    )
                )
                if bidi_removed:
                    layout_signals.append("bidirectional_controls_present")
                if recovered:
                    layout_signals.append("hyperlinks_recovered")
                continue
            if isinstance(item, Table):
                for row in item.rows:
                    raw_values: list[str] = []
                    recovered = False
                    for cell in row.cells:
                        cell_value, cell_recovered = _docx_value_with_hyperlinks(
                            cell.text,
                            *cell.paragraphs,
                        )
                        raw_values.append(cell_value.replace("\x00", "").strip())
                        recovered = recovered or cell_recovered
                    raw_value = " | ".join(raw_values)
                    value, bidi_removed = _sanitize_extracted_text(raw_value)
                    if bidi_removed:
                        layout_signals.append("bidirectional_controls_present")
                    if recovered:
                        layout_signals.append("hyperlinks_recovered")
                    if value.strip(" |"):
                        values.append(("table", value))
        if any(
            paragraph.text.strip()
            for section in document.sections
            for part in (section.header, section.footer)
            for paragraph in part.paragraphs
        ):
            layout_signals.append("header_footer_present")
    except Exception as exc:
        raise UnsafeDocument("malformed_docx") from exc
    values = [(kind, value) for kind, value in values if value]
    paragraphs = [value for _, value in values]
    plain_text = "\n".join(paragraphs)
    if len(plain_text) > limits.max_extracted_characters:
        raise UnsafeDocument("extracted_text_limit_exceeded")
    blocks = _blocks_from_docx_values(values, limits.max_extracted_blocks)
    return ExtractionResult(
        plain_text=plain_text,
        reading_order=blocks,
        page_count=1,
        image_only=len(plain_text.strip()) < 40,
        warnings=("limited_extractable_text",) if len(plain_text.strip()) < 40 else (),
        parser_version=PARSER_VERSION,
        layout_signals=tuple(dict.fromkeys(layout_signals)),
    )


def _blocks_from_pages(page_texts: Iterable[str], max_blocks: int) -> tuple[ExtractedBlock, ...]:
    blocks: list[ExtractedBlock] = []
    global_offset = 0
    for page_number, page_text in enumerate(page_texts, start=1):
        cursor = 0
        for raw_line in page_text.splitlines():
            text = raw_line.strip()
            if not text:
                cursor += len(raw_line) + 1
                continue
            relative = page_text.find(text, cursor)
            relative = cursor if relative < 0 else relative
            cursor = relative + len(text)
            kind = "bullet" if _BULLET_PREFIX.match(text) else _guess_block_kind(text)
            if len(blocks) >= max_blocks:
                raise UnsafeDocument("extracted_block_limit_exceeded")
            blocks.append(
                ExtractedBlock(
                    kind=kind,
                    text=text,
                    confidence_basis_points=9_000,
                    spans=(
                        SourceSpanView(
                            page=page_number,
                            start=global_offset + relative,
                            end=global_offset + relative + len(text),
                        ),
                    ),
                )
            )
        global_offset += len(page_text) + 3
    return tuple(blocks)


def _blocks_from_docx_values(
    values: list[tuple[str, str]],
    max_blocks: int,
) -> tuple[ExtractedBlock, ...]:
    blocks: list[ExtractedBlock] = []
    offset = 0
    for source_kind, value in values:
        if len(blocks) >= max_blocks:
            raise UnsafeDocument("extracted_block_limit_exceeded")
        kind = (
            source_kind
            if source_kind in {"bullet", "table"}
            else ("bullet" if _BULLET_PREFIX.match(value) else _guess_block_kind(value))
        )
        blocks.append(
            ExtractedBlock(
                kind=kind,
                text=value,
                confidence_basis_points=8_500 if kind == "table" else 9_000,
                spans=(SourceSpanView(page=1, start=offset, end=offset + len(value)),),
            )
        )
        offset += len(value) + 1
    return tuple(blocks)


def _guess_block_kind(text: str) -> str:
    if len(text) <= 64 and (text.isupper() or text.endswith(":")):
        return "heading"
    return "paragraph"


def _is_docx_list_paragraph(paragraph: Paragraph) -> bool:
    """Recognize Word list items even when their marker is formatting-only."""
    paragraph_properties = paragraph._element.pPr
    if paragraph_properties is not None and paragraph_properties.numPr is not None:
        return True
    style = paragraph.style
    while style is not None:
        style_name = (style.name or "").strip().casefold()
        if style_name.startswith("list") or "bullet" in style_name:
            return True
        style = style.base_style
    return False


def stable_block_id(document_id: UUID, block: ExtractedBlock, index: int) -> UUID:
    span = block.spans[0] if block.spans else None
    material = f"{document_id}:{index}:{span.page if span else 0}:{span.start if span else 0}"
    return uuid5(_BLOCK_NAMESPACE, material)


def normalize_extracted_text(value: str) -> str:
    """Repair common PDF/DOCX extraction artifacts in contact and link lines."""

    normalized = _CONTACT_ICON_GLYPHS.sub("", value)
    normalized = _PHONE_LABEL_GLUE.sub("", normalized)
    normalized = _DUPLICATE_PROFILE_HOST.sub(r"\1.com", normalized)
    normalized = _SLASH_BEFORE_PROFILE_HOST.sub("", normalized)
    normalized = re.sub(r"[ \t]{2,}", " ", normalized)
    normalized = re.sub(r"\s+\|", " |", normalized)
    return normalized.strip()


def _validate_docx_relationships(value: bytes) -> None:
    """Allow only safe web hyperlinks among external DOCX relationships."""

    for relationship in _DOCX_RELATIONSHIP_TAG.finditer(value):
        try:
            attributes = {
                name.decode("ascii").casefold(): html.unescape(target.decode("utf-8"))
                for name, _quote, target in _DOCX_XML_ATTRIBUTE.findall(relationship.group())
            }
        except UnicodeDecodeError as exc:
            raise UnsafeDocument("malformed_docx") from exc
        if attributes.get("targetmode", "").casefold() != "external":
            continue
        relationship_type = attributes.get("type", "")
        target = attributes.get("target", "")
        if not relationship_type.casefold().endswith("/hyperlink"):
            raise UnsafeDocument("embedded_object_rejected")
        if safe_http_hyperlink(target) is None:
            raise UnsafeDocument("embedded_object_rejected")


def _docx_value_with_hyperlinks(value: str, *paragraphs: Paragraph) -> tuple[str, bool]:
    targets: list[str] = []
    for paragraph in paragraphs:
        for item in paragraph.iter_inner_content():
            if isinstance(item, Hyperlink):
                target = safe_http_hyperlink(item.url)
                if target is not None:
                    targets.append(target)
    return append_hyperlink_targets(value, tuple(targets))


def _sanitize_extracted_text(value: str) -> tuple[str, bool]:
    without_bidi = value.translate(_BIDI_CONTROLS)
    sanitized = normalize_extracted_text(without_bidi)
    return sanitized, sanitized != value or without_bidi != value
