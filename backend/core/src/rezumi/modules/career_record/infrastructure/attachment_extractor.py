"""Bounded, payload-free PDF/DOCX validation for evidence attachments."""

import asyncio
import stat
from collections.abc import Iterable
from pathlib import Path, PurePosixPath
from typing import Never
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from rezumi.modules.career_record.application.attachment_workflow import (
    AttachmentExtractionSummary,
    AttachmentLimits,
    AttachmentMediaType,
    SafeAttachmentError,
    UnsafeAttachment,
)

ATTACHMENT_PARSER_VERSION = "rezumi-attachment-parser/1.1.0"


class BoundedAttachmentExtractor:
    """Validate hostile inputs and retain counts only, never extracted text."""

    async def extract(
        self, path: Path, media_type: AttachmentMediaType, limits: AttachmentLimits
    ) -> AttachmentExtractionSummary:
        try:
            return await asyncio.wait_for(
                asyncio.to_thread(extract_attachment_document, path, media_type, limits),
                limits.processing_timeout_seconds,
            )
        except TimeoutError as exc:
            raise RuntimeError("attachment extraction timed out") from exc


def extract_attachment_document(
    path: Path,
    media_type: AttachmentMediaType | str,
    limits: AttachmentLimits,
) -> AttachmentExtractionSummary:
    """Run bounded attachment validation inside the caller's isolation boundary."""
    resolved_type = (
        media_type
        if isinstance(media_type, AttachmentMediaType)
        else AttachmentMediaType(media_type)
    )
    if not path.is_file() or path.is_symlink() or path.stat().st_size > limits.max_upload_bytes:
        _unsafe()
    if resolved_type is AttachmentMediaType.PDF:
        return _pdf_summary(path, limits)
    if resolved_type is AttachmentMediaType.DOCX:
        return _docx_summary(path, limits)
    _unsafe()


def _pdf_summary(path: Path, limits: AttachmentLimits) -> AttachmentExtractionSummary:
    raw = path.read_bytes()
    if not raw.startswith(b"%PDF-") or b"PK\x03\x04" in raw:
        _unsafe()
    if any(token in raw for token in (b"/JavaScript", b"/JS", b"/Launch", b"/EmbeddedFile")):
        _unsafe()
    if b"%%EOF" not in raw[-2_048:]:
        _unsafe()
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted or not 1 <= len(reader.pages) <= limits.max_pdf_pages:
            _unsafe()
        page_texts = tuple((page.extract_text() or "").replace("\x00", "") for page in reader.pages)
    except UnsafeAttachment:
        raise
    except Exception as exc:
        raise UnsafeAttachment(SafeAttachmentError.INVALID_DOCUMENT_STRUCTURE) from exc
    characters, blocks = _text_counts(page_texts, limits)
    return AttachmentExtractionSummary(
        format_valid=True,
        page_count=len(page_texts),
        extracted_characters=characters,
        extracted_blocks=blocks,
        archive_entries=0,
        archive_uncompressed_bytes=0,
        max_archive_ratio=0,
        parser_version=ATTACHMENT_PARSER_VERSION,
    )


def _docx_summary(path: Path, limits: AttachmentLimits) -> AttachmentExtractionSummary:
    raw_prefix = path.read_bytes()[:1_024]
    if not raw_prefix.startswith(b"PK\x03\x04") or b"%PDF-" in raw_prefix:
        _unsafe()
    archive_entries = 0
    total_uncompressed = 0
    observed_ratio = 0
    try:
        with ZipFile(path) as archive:
            infos = archive.infolist()
            archive_entries = len(infos)
            if not 1 <= archive_entries <= limits.max_archive_entries:
                _unsafe()
            names: set[str] = set()
            total_compressed = 0
            for info in infos:
                normalized = PurePosixPath(info.filename.replace("\\", "/"))
                if normalized.is_absolute() or ".." in normalized.parts:
                    _unsafe()
                mode = info.external_attr >> 16
                if mode and stat.S_ISLNK(mode):
                    _unsafe()
                if info.flag_bits & 0x1:
                    _unsafe()
                total_uncompressed += info.file_size
                compressed = max(1, info.compress_size)
                total_compressed += compressed
                observed_ratio = max(
                    observed_ratio,
                    (info.file_size + compressed - 1) // compressed,
                )
                names.add(info.filename.casefold())
            total_ratio = (total_uncompressed + total_compressed - 1) // max(1, total_compressed)
            observed_ratio = max(observed_ratio, total_ratio)
            if (
                total_uncompressed > limits.max_archive_uncompressed_bytes
                or observed_ratio > limits.max_archive_ratio
            ):
                _unsafe()
            if "[content_types].xml" not in names or "word/document.xml" not in names:
                _unsafe()
            if any(
                name == "word/vbaproject.bin"
                or name.startswith("word/activex/")
                or name.startswith("word/embeddings/")
                for name in names
            ):
                _unsafe()
            for info in infos:
                if info.filename.casefold().endswith(".rels"):
                    relationships = archive.read(info)
                    lowered = relationships.lower()
                    if b'targetmode="external"' in lowered or b"targetmode='external'" in lowered:
                        _unsafe()
            if archive.testzip() is not None:
                _unsafe()
    except UnsafeAttachment:
        raise
    except (BadZipFile, OSError, KeyError) as exc:
        raise UnsafeAttachment(SafeAttachmentError.INVALID_DOCUMENT_STRUCTURE) from exc

    try:
        document = Document(str(path))
        values = [paragraph.text.replace("\x00", "") for paragraph in document.paragraphs]
        values.extend(
            " | ".join(cell.text.replace("\x00", "") for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
    except Exception as exc:
        raise UnsafeAttachment(SafeAttachmentError.INVALID_DOCUMENT_STRUCTURE) from exc
    characters, blocks = _text_counts(values, limits)
    return AttachmentExtractionSummary(
        format_valid=True,
        page_count=1,
        extracted_characters=characters,
        extracted_blocks=blocks,
        archive_entries=archive_entries,
        archive_uncompressed_bytes=total_uncompressed,
        max_archive_ratio=observed_ratio,
        parser_version=ATTACHMENT_PARSER_VERSION,
    )


def _text_counts(values: Iterable[str], limits: AttachmentLimits) -> tuple[int, int]:
    characters = 0
    blocks = 0
    for value in values:
        characters += len(value)
        blocks += sum(bool(line.strip()) for line in value.splitlines())
        if characters > limits.max_extracted_characters or blocks > limits.max_extracted_blocks:
            _unsafe()
    return characters, blocks


def _unsafe() -> Never:
    raise UnsafeAttachment(SafeAttachmentError.INVALID_DOCUMENT_STRUCTURE)
