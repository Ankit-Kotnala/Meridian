"""Hostile PDF/DOCX admission and local extraction tests."""

import hashlib
import json
import time
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter

from rezumi.modules.resume_health.application.models import DocumentLimits
from rezumi.modules.resume_health.domain import ResumeMediaType
from rezumi.modules.resume_health.domain.errors import UnsafeDocument
from rezumi.modules.resume_health.infrastructure.extractors import (
    PARSER_VERSION,
    LocalDocumentExtractor,
    LocalDocumentTextExtractor,
    normalize_extracted_text,
)
from rezumi.modules.resume_health.infrastructure.isolated_extractor import (
    IsolatedDocumentExtractor,
)
from rezumi.modules.resume_health.infrastructure.layout import (
    LAYOUT_ANALYZER_VERSION,
    LocalLayoutAnalyzer,
)

FIXTURES = Path(__file__).resolve().parents[4] / "frontend" / "test-fixtures" / "generated"


def test_committed_fixture_manifest_matches_bounded_fictional_corpus() -> None:
    manifest = json.loads((FIXTURES / "manifest.json").read_text(encoding="utf-8"))

    assert set(manifest) == {
        "adversarial-layout.docx",
        "fictional-resume.docx",
        "fictional-resume.pdf",
        "image-only.pdf",
        "long-resume.docx",
        "two-column.pdf",
    }
    for filename, expected in manifest.items():
        value = (FIXTURES / filename).read_bytes()
        assert len(value) == expected["bytes"]
        assert hashlib.sha256(value).hexdigest() == expected["sha256"]


@pytest.mark.asyncio
async def test_wrong_signature_fails_before_parser(tmp_path: Path) -> None:
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"not a PDF")

    with pytest.raises(UnsafeDocument, match="document_signature_mismatch"):
        await LocalDocumentExtractor().extract(path, ResumeMediaType.PDF.value, DocumentLimits())


@pytest.mark.asyncio
async def test_docx_macro_traversal_and_expansion_fail_closed(tmp_path: Path) -> None:
    cases = [
        (
            "macro.docx",
            [
                ("[Content_Types].xml", b"x"),
                ("word/document.xml", b"x"),
                ("word/vbaProject.bin", b"x"),
            ],
            "macro_document_rejected",
        ),
        (
            "traversal.docx",
            [("[Content_Types].xml", b"x"), ("word/document.xml", b"x"), ("../escape", b"x")],
            "archive_path_traversal",
        ),
    ]
    for filename, members, code in cases:
        path = tmp_path / filename
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            for name, value in members:
                archive.writestr(name, value)
        with pytest.raises(UnsafeDocument, match=code):
            await LocalDocumentExtractor().extract(
                path, ResumeMediaType.DOCX.value, DocumentLimits()
            )


@pytest.mark.asyncio
async def test_clean_docx_extracts_text_blocks_and_spans(tmp_path: Path) -> None:
    path = tmp_path / "fictional-resume.docx"
    document = Document()
    document.add_heading("Experience", level=1)
    document.add_paragraph(
        "Built a fictional service and improved delivery quality.", style="List Bullet"
    )
    document.add_heading("Education", level=1)
    document.add_paragraph("Fictional University, 2024")
    document.save(path)

    result = await LocalDocumentExtractor().extract(
        path, ResumeMediaType.DOCX.value, DocumentLimits()
    )

    assert "fictional service" in result.plain_text
    assert result.reading_order
    assert any(
        block.kind == "bullet" and "fictional service" in block.text
        for block in result.reading_order
    )
    assert all(block.spans and block.spans[0].page == 1 for block in result.reading_order)
    assert not result.image_only


@pytest.mark.asyncio
async def test_docx_paragraphs_and_tables_keep_source_order(tmp_path: Path) -> None:
    path = tmp_path / "interleaved.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Middle role"
    table.cell(0, 1).text = "2024"
    document.add_paragraph("After table")
    document.save(path)

    result = await LocalDocumentExtractor().extract(
        path,
        ResumeMediaType.DOCX.value,
        DocumentLimits(),
    )

    assert [block.text for block in result.reading_order] == [
        "Before table",
        "Middle role | 2024",
        "After table",
    ]
    assert result.plain_text.splitlines() == [block.text for block in result.reading_order]


@pytest.mark.asyncio
async def test_table_heavy_docx_is_preserved_and_flagged_for_layout_review(
    tmp_path: Path,
) -> None:
    path = tmp_path / "table-heavy.docx"
    document = Document()
    for index in range(3):
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = f"Fictional role {index}"
        table.cell(0, 1).text = "2024"
    document.save(path)

    result = await LocalDocumentExtractor().extract(
        path,
        ResumeMediaType.DOCX.value,
        DocumentLimits(),
    )

    assert [block.kind for block in result.reading_order] == ["table", "table", "table"]
    assert "table_heavy_layout" in result.warnings

    text_only = await LocalDocumentTextExtractor().extract_text(
        path,
        ResumeMediaType.DOCX.value,
        DocumentLimits(),
    )
    analyzed = await LocalLayoutAnalyzer().analyze(
        path,
        ResumeMediaType.DOCX.value,
        text_only,
        DocumentLimits(),
    )
    assert "table_heavy_layout" not in text_only.warnings
    assert "table_heavy_layout" in analyzed.warnings


@pytest.mark.asyncio
async def test_header_footer_unusual_font_and_bidi_controls_are_handled_safely(
    tmp_path: Path,
) -> None:
    path = tmp_path / "hostile-layout.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "FICTIONAL HEADER"
    document.sections[0].footer.paragraphs[0].text = "FICTIONAL FOOTER"
    paragraph = document.add_paragraph("Experience")
    paragraph.runs[0].font.name = "Papyrus"
    document.add_paragraph("Principal\u202e Engineer | Fictional Labs | 2024")
    document.save(path)

    result = await LocalDocumentExtractor().extract(
        path,
        ResumeMediaType.DOCX.value,
        DocumentLimits(),
    )

    assert "\u202e" not in result.plain_text
    assert "FICTIONAL HEADER" not in result.plain_text
    assert "FICTIONAL FOOTER" not in result.plain_text
    assert {"header_footer_excluded", "bidirectional_controls_removed"} <= set(result.warnings)


@pytest.mark.asyncio
async def test_long_document_fails_at_character_boundary(tmp_path: Path) -> None:
    path = tmp_path / "long.docx"
    document = Document()
    for index in range(60):
        document.add_paragraph(f"Fictional achievement {index} " + "x" * 80)
    document.save(path)

    with pytest.raises(UnsafeDocument, match="extracted_text_limit_exceeded"):
        await LocalDocumentExtractor().extract(
            path,
            ResumeMediaType.DOCX.value,
            DocumentLimits(max_extracted_characters=1_000),
        )


@pytest.mark.asyncio
async def test_committed_clean_pdf_and_image_only_pdf_are_classified() -> None:
    extractor = LocalDocumentExtractor()
    clean = await extractor.extract(
        FIXTURES / "fictional-resume.pdf", ResumeMediaType.PDF.value, DocumentLimits()
    )
    image_only = await extractor.extract(
        FIXTURES / "image-only.pdf", ResumeMediaType.PDF.value, DocumentLimits()
    )

    assert "ALEX RIVERA" in clean.plain_text
    assert clean.reading_order
    assert not clean.image_only
    assert image_only.image_only
    assert image_only.warnings == ("image_only_pdf",)


@pytest.mark.asyncio
async def test_committed_adversarial_layout_corpus_has_expected_safe_signals() -> None:
    extractor = LocalDocumentExtractor()
    two_column = await extractor.extract(
        FIXTURES / "two-column.pdf",
        ResumeMediaType.PDF.value,
        DocumentLimits(),
    )
    adversarial = await extractor.extract(
        FIXTURES / "adversarial-layout.docx",
        ResumeMediaType.DOCX.value,
        DocumentLimits(),
    )

    assert "multi_column_layout" in two_column.warnings
    assert "reading_order_uncertain" not in two_column.warnings
    assert two_column.plain_text.index("EXPERIENCE") < two_column.plain_text.index("SKILLS")
    assert two_column.plain_text.index("Engineer | Fictional Alpha") < two_column.plain_text.index(
        "Python, product strategy"
    )
    assert {
        "bidirectional_controls_removed",
        "header_footer_excluded",
        "table_heavy_layout",
    } <= set(adversarial.warnings)
    assert "\u202e" not in adversarial.plain_text
    assert "FICTIONAL REPEATED HEADER" not in adversarial.plain_text

    long_resume = await extractor.extract(
        FIXTURES / "long-resume.docx",
        ResumeMediaType.DOCX.value,
        DocumentLimits(),
    )
    assert len(long_resume.reading_order) == 601
    assert len(long_resume.plain_text) < DocumentLimits().max_extracted_characters


@pytest.mark.asyncio
async def test_encrypted_malformed_and_polyglot_pdf_fail_closed(tmp_path: Path) -> None:
    encrypted = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("test-only-password")
    with encrypted.open("wb") as output:
        writer.write(output)
    malformed = tmp_path / "malformed.pdf"
    malformed.write_bytes(b"%PDF-1.7\ntruncated")
    polyglot = tmp_path / "polyglot.pdf"
    polyglot.write_bytes(
        (FIXTURES / "fictional-resume.pdf").read_bytes() + b"PK\x03\x04embedded-archive"
    )

    for path, code in (
        (encrypted, "encrypted_document"),
        (malformed, "malformed_pdf"),
        (polyglot, "polyglot_document_rejected"),
    ):
        with pytest.raises(UnsafeDocument, match=code):
            await LocalDocumentExtractor().extract(
                path, ResumeMediaType.PDF.value, DocumentLimits()
            )


@pytest.mark.asyncio
async def test_docx_entry_expansion_and_ratio_limits_fail_before_open(tmp_path: Path) -> None:
    path = tmp_path / "bounded.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"x" * 100)
        archive.writestr("word/document.xml", b"y" * 100)
        archive.writestr("word/styles.xml", b"z" * 100)

    cases = (
        (DocumentLimits(max_archive_entries=2), "archive_entry_limit_exceeded"),
        (
            DocumentLimits(max_archive_uncompressed_bytes=150),
            "archive_expansion_limit_exceeded",
        ),
        (DocumentLimits(max_archive_ratio=1), "archive_ratio_limit_exceeded"),
    )
    for limits, code in cases:
        with pytest.raises(UnsafeDocument, match=code):
            await LocalDocumentExtractor().extract(path, ResumeMediaType.DOCX.value, limits)


@pytest.mark.asyncio
async def test_extraction_timeout_returns_safe_failure(tmp_path: Path) -> None:
    class SlowExtractor(LocalDocumentExtractor):
        def _extract(self, path: Path, media_type: str, limits: DocumentLimits) -> object:
            del path, media_type, limits
            time.sleep(0.05)
            raise AssertionError("late parser result must be ignored")

    path = tmp_path / "slow.pdf"
    path.write_bytes(b"%PDF-1.7\n%%EOF")
    with pytest.raises(UnsafeDocument, match="document_processing_timeout"):
        await SlowExtractor().extract(
            path,
            ResumeMediaType.PDF.value,
            DocumentLimits(processing_timeout_seconds=0.001),
        )


@pytest.mark.asyncio
async def test_isolated_extractor_returns_validated_child_result(tmp_path: Path) -> None:
    source = tmp_path / "fictional-resume.pdf"
    source.write_bytes((FIXTURES / "fictional-resume.pdf").read_bytes())

    result = await IsolatedDocumentExtractor().extract(
        source,
        ResumeMediaType.PDF.value,
        DocumentLimits(temp_root=tmp_path),
    )

    assert "ALEX RIVERA" in result.plain_text
    assert result.parser_version == f"{PARSER_VERSION}+{LAYOUT_ANALYZER_VERSION}"
    assert not list(tmp_path.glob("parser-*"))


@pytest.mark.asyncio
async def test_isolated_extractor_kills_timed_out_child_and_cleans_workspace(
    tmp_path: Path,
) -> None:
    source = tmp_path / "fictional-resume.pdf"
    source.write_bytes((FIXTURES / "fictional-resume.pdf").read_bytes())

    with pytest.raises(UnsafeDocument, match="document_processing_timeout"):
        await IsolatedDocumentExtractor().extract(
            source,
            ResumeMediaType.PDF.value,
            DocumentLimits(
                processing_timeout_seconds=0.000_001,
                temp_root=tmp_path,
            ),
        )

    assert not list(tmp_path.glob("parser-*"))


def _pdf_pages(*streams: str) -> bytes:
    encoded = [value.encode("ascii") for value in streams]
    page_count = len(encoded)
    page_ids = list(range(3, 3 + page_count))
    font_id = 3 + page_count
    content_ids = list(range(font_id + 1, font_id + 1 + page_count))
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode("ascii"),
    ]
    for content_id in content_ids:
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_id} 0 R >>"
            ).encode("ascii")
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for stream in encoded:
        objects.append(
            b"<< /Length "
            + str(len(stream)).encode("ascii")
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n").encode(
            "ascii"
        )
    )
    return bytes(output)


@pytest.mark.asyncio
async def test_row_wise_two_column_pdf_is_read_left_column_then_right(tmp_path: Path) -> None:
    path = tmp_path / "interleaved-columns.pdf"
    path.write_bytes(
        _pdf_pages(
            "BT\n/F1 10 Tf\n40 760 Td\n(LEFT-A) Tj\nET\n"
            "BT\n/F1 10 Tf\n320 760 Td\n(RIGHT-A) Tj\nET\n"
            "BT\n/F1 10 Tf\n40 740 Td\n(LEFT-B) Tj\nET\n"
            "BT\n/F1 10 Tf\n320 740 Td\n(RIGHT-B) Tj\nET"
        )
    )

    result = await LocalDocumentExtractor().extract(
        path, ResumeMediaType.PDF.value, DocumentLimits()
    )

    assert result.plain_text.splitlines() == ["LEFT-A", "LEFT-B", "RIGHT-A", "RIGHT-B"]
    assert "multi_column_layout" in result.warnings
    assert "reading_order_uncertain" not in result.warnings


@pytest.mark.asyncio
async def test_repeating_pdf_header_is_excluded_from_career_text(tmp_path: Path) -> None:
    path = tmp_path / "running-header.pdf"
    path.write_bytes(
        _pdf_pages(
            "BT\n/F1 10 Tf\n48 770 Td\n(CONFIDENTIAL) Tj\n0 -40 Td\n(ALPHA BODY) Tj\nET",
            "BT\n/F1 10 Tf\n48 770 Td\n(CONFIDENTIAL) Tj\n0 -40 Td\n(BETA BODY) Tj\nET",
        )
    )

    result = await LocalDocumentExtractor().extract(
        path, ResumeMediaType.PDF.value, DocumentLimits()
    )

    assert "ALPHA BODY" in result.plain_text
    assert "BETA BODY" in result.plain_text
    assert "CONFIDENTIAL" not in result.plain_text
    assert "header_footer_excluded" in result.warnings


def test_normalize_extracted_text_repairs_pdf_contact_link_artifacts() -> None:
    raw = (
        "♂phone+91 88001 45975 | ✉ankit.kotnala12@gmail.com | "
        "/linkedinlinkedin.com/in/ankit-kotnala- | /githubgithub.com/Ankit-Kotnala"
    )

    assert normalize_extracted_text(raw) == (
        "+91 88001 45975 | ankit.kotnala12@gmail.com | "
        "linkedin.com/in/ankit-kotnala- | github.com/Ankit-Kotnala"
    )


@pytest.mark.asyncio
async def test_pdf_javascript_and_docx_embeddings_fail_closed(tmp_path: Path) -> None:
    pdf_path = tmp_path / "active.pdf"
    pdf_path.write_bytes((FIXTURES / "fictional-resume.pdf").read_bytes() + b"\n/JavaScript\n")
    with pytest.raises(UnsafeDocument, match="active_content_rejected"):
        await LocalDocumentExtractor().extract(
            pdf_path, ResumeMediaType.PDF.value, DocumentLimits()
        )

    docx_path = tmp_path / "embedded.docx"
    document = Document()
    document.add_paragraph("Fictional evidence")
    document.save(docx_path)
    with ZipFile(docx_path, "a") as archive:
        archive.writestr("word/embeddings/oleObject1.bin", b"x")
    with pytest.raises(UnsafeDocument, match="embedded_object_rejected"):
        await LocalDocumentExtractor().extract(
            docx_path, ResumeMediaType.DOCX.value, DocumentLimits()
        )
