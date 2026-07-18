"""Hostile PDF/DOCX admission and local extraction tests."""

import time
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter

from careeros.modules.resume_health.application.models import DocumentLimits
from careeros.modules.resume_health.domain import ResumeMediaType
from careeros.modules.resume_health.domain.errors import UnsafeDocument
from careeros.modules.resume_health.infrastructure.extractors import LocalDocumentExtractor

FIXTURES = Path(__file__).resolve().parents[3] / "test-fixtures" / "generated"


@pytest.mark.asyncio
async def test_wrong_signature_fails_before_parser(tmp_path: Path) -> None:
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"not a PDF")

    with pytest.raises(UnsafeDocument, match="document_signature_mismatch"):
        await LocalDocumentExtractor().extract(path, ResumeMediaType.PDF.value, DocumentLimits())


@pytest.mark.asyncio
async def test_docx_macro_traversal_and_expansion_fail_closed(tmp_path: Path) -> None:
    cases = [
        (
            "macro.docx",
            [
                ("[Content_Types].xml", b"x"),
                ("word/document.xml", b"x"),
                ("word/vbaProject.bin", b"x"),
            ],
            "macro_document_rejected",
        ),
        (
            "traversal.docx",
            [("[Content_Types].xml", b"x"), ("word/document.xml", b"x"), ("../escape", b"x")],
            "archive_path_traversal",
        ),
    ]
    for filename, members, code in cases:
        path = tmp_path / filename
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            for name, value in members:
                archive.writestr(name, value)
        with pytest.raises(UnsafeDocument, match=code):
            await LocalDocumentExtractor().extract(
                path, ResumeMediaType.DOCX.value, DocumentLimits()
            )


@pytest.mark.asyncio
async def test_clean_docx_extracts_text_blocks_and_spans(tmp_path: Path) -> None:
    path = tmp_path / "fictional-resume.docx"
    document = Document()
    document.add_heading("Experience", level=1)
    document.add_paragraph(
        "Built a fictional service and improved delivery quality.", style="List Bullet"
    )
    document.add_heading("Education", level=1)
    document.add_paragraph("Fictional University, 2024")
    document.save(path)

    result = await LocalDocumentExtractor().extract(
        path, ResumeMediaType.DOCX.value, DocumentLimits()
    )

    assert "fictional service" in result.plain_text
    assert result.reading_order
    assert all(block.spans and block.spans[0].page == 1 for block in result.reading_order)
    assert not result.image_only


@pytest.mark.asyncio
async def test_committed_clean_pdf_and_image_only_pdf_are_classified() -> None:
    extractor = LocalDocumentExtractor()
    clean = await extractor.extract(
        FIXTURES / "fictional-resume.pdf", ResumeMediaType.PDF.value, DocumentLimits()
    )
    image_only = await extractor.extract(
        FIXTURES / "image-only.pdf", ResumeMediaType.PDF.value, DocumentLimits()
    )

    assert "ALEX RIVERA" in clean.plain_text
    assert clean.reading_order
    assert not clean.image_only
    assert image_only.image_only
    assert image_only.warnings == ("image_only_pdf",)


@pytest.mark.asyncio
async def test_encrypted_malformed_and_polyglot_pdf_fail_closed(tmp_path: Path) -> None:
    encrypted = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("test-only-password")
    with encrypted.open("wb") as output:
        writer.write(output)
    malformed = tmp_path / "malformed.pdf"
    malformed.write_bytes(b"%PDF-1.7\ntruncated")
    polyglot = tmp_path / "polyglot.pdf"
    polyglot.write_bytes(
        (FIXTURES / "fictional-resume.pdf").read_bytes() + b"PK\x03\x04embedded-archive"
    )

    for path, code in (
        (encrypted, "encrypted_document"),
        (malformed, "malformed_pdf"),
        (polyglot, "polyglot_document_rejected"),
    ):
        with pytest.raises(UnsafeDocument, match=code):
            await LocalDocumentExtractor().extract(
                path, ResumeMediaType.PDF.value, DocumentLimits()
            )


@pytest.mark.asyncio
async def test_docx_entry_expansion_and_ratio_limits_fail_before_open(tmp_path: Path) -> None:
    path = tmp_path / "bounded.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"x" * 100)
        archive.writestr("word/document.xml", b"y" * 100)
        archive.writestr("word/styles.xml", b"z" * 100)

    cases = (
        (DocumentLimits(max_archive_entries=2), "archive_entry_limit_exceeded"),
        (
            DocumentLimits(max_archive_uncompressed_bytes=150),
            "archive_expansion_limit_exceeded",
        ),
        (DocumentLimits(max_archive_ratio=1), "archive_ratio_limit_exceeded"),
    )
    for limits, code in cases:
        with pytest.raises(UnsafeDocument, match=code):
            await LocalDocumentExtractor().extract(path, ResumeMediaType.DOCX.value, limits)


@pytest.mark.asyncio
async def test_extraction_timeout_returns_safe_failure(tmp_path: Path) -> None:
    class SlowExtractor(LocalDocumentExtractor):
        def _extract(self, path: Path, media_type: str, limits: DocumentLimits) -> object:
            del path, media_type, limits
            time.sleep(0.05)
            raise AssertionError("late parser result must be ignored")

    path = tmp_path / "slow.pdf"
    path.write_bytes(b"%PDF-1.7\n%%EOF")
    with pytest.raises(UnsafeDocument, match="document_processing_timeout"):
        await SlowExtractor().extract(
            path,
            ResumeMediaType.PDF.value,
            DocumentLimits(processing_timeout_seconds=0.001),
        )
