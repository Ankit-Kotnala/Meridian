"""Credential stripping and killable attachment-parser regressions."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from careeros.foundation.sandbox import sanitized_parser_environment
from careeros.modules.career_record.application.attachment_workflow import (
    AttachmentLimits,
    AttachmentMediaType,
    SafeAttachmentError,
    UnsafeAttachment,
)
from careeros.modules.career_record.infrastructure import IsolatedAttachmentExtractor


def test_parser_environment_contains_no_parent_credentials(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("CAREEROS_DATABASE_URL", "fictional-sensitive-value")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "fictional-sensitive-value")

    child = sanitized_parser_environment(tmp_path)

    assert not any(key.startswith("CAREEROS_") for key in child)
    assert "AWS_SECRET_ACCESS_KEY" not in child
    assert child["PYTHONNOUSERSITE"] == "1"
    assert child["TMPDIR"] == str(tmp_path.resolve())


def test_parser_audit_guard_denies_socket_creation() -> None:
    script = (
        "import socket\n"
        "from careeros.foundation.sandbox import install_parser_egress_guard\n"
        "install_parser_egress_guard()\n"
        "try:\n"
        "    socket.socket()\n"
        "except PermissionError:\n"
        "    raise SystemExit(0)\n"
        "raise SystemExit(1)\n"
    )

    completed = subprocess.run(  # noqa: S603 - fixed interpreter and test script.
        [sys.executable, "-c", script],
        check=False,
        capture_output=True,
        text=True,
        timeout=10,
    )

    assert completed.returncode == 0


@pytest.mark.asyncio
async def test_isolated_attachment_extractor_returns_counts_and_cleans_child(
    tmp_path: Path,
) -> None:
    path = tmp_path / "evidence.docx"
    document = Document()
    document.add_heading("Fictional Evidence", level=1)
    document.add_paragraph("A deliberately fictional supporting statement.")
    document.save(path)

    result = await IsolatedAttachmentExtractor().extract(
        path,
        AttachmentMediaType.DOCX,
        AttachmentLimits(temp_root=tmp_path.resolve()),
    )

    assert result.format_valid
    assert result.extracted_characters > 0
    assert result.extracted_blocks == 2
    assert not list(tmp_path.glob("attachment-parser-*"))


@pytest.mark.asyncio
async def test_isolated_attachment_extractor_maps_invalid_bytes_and_timeout(
    tmp_path: Path,
) -> None:
    invalid = tmp_path / "invalid.pdf"
    invalid.write_bytes(b"not a PDF")
    with pytest.raises(UnsafeAttachment) as raised:
        await IsolatedAttachmentExtractor().extract(
            invalid,
            AttachmentMediaType.PDF,
            AttachmentLimits(temp_root=tmp_path.resolve()),
        )
    assert raised.value.code is SafeAttachmentError.INVALID_DOCUMENT_STRUCTURE

    with pytest.raises(RuntimeError, match="timed out"):
        await IsolatedAttachmentExtractor().extract(
            invalid,
            AttachmentMediaType.PDF,
            AttachmentLimits(
                processing_timeout_seconds=0.000_001,
                temp_root=tmp_path.resolve(),
            ),
        )
    assert not list(tmp_path.glob("attachment-parser-*"))
