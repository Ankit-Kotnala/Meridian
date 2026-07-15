"""Provider-adapter tests for private Career Record attachments."""

from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from careeros.modules.career_record.application.attachment_workflow import (
    AttachmentDownloadPurpose,
    AttachmentLimits,
    AttachmentMediaType,
    SafeAttachmentError,
    ScanVerdict,
    UnsafeAttachment,
)
from careeros.modules.career_record.infrastructure.attachment_extractor import (
    BoundedAttachmentExtractor,
)
from careeros.modules.career_record.infrastructure.attachment_scanner import (
    AttachmentClamAvOptions,
    AttachmentClamAvScanner,
)
from careeros.modules.career_record.infrastructure.attachment_storage import (
    AttachmentS3ObjectStorage,
    AttachmentS3Options,
)


class _ClosingBody:
    def __init__(self, value: bytes) -> None:
        self._stream = BytesIO(value)
        self.closed = False

    def read(self, length: int = -1) -> bytes:
        return self._stream.read(length)

    def close(self) -> None:
        self.closed = True
        self._stream.close()


class _ObjectClient:
    def __init__(self, bodies: list[_ClosingBody]) -> None:
        self._bodies = iter(bodies)

    def get_object(self, **kwargs: Any) -> dict[str, _ClosingBody]:
        del kwargs
        return {"Body": next(self._bodies)}


class _Signer:
    def __init__(self) -> None:
        self.calls: list[tuple[str, dict[str, Any]]] = []

    def generate_presigned_url(self, operation: str, **kwargs: Any) -> str:
        self.calls.append((operation, kwargs))
        return f"https://objects.invalid/{operation}/{len(self.calls)}"


def _options() -> AttachmentS3Options:
    return AttachmentS3Options(
        internal_endpoint_url="http://s3.invalid",
        public_endpoint_url="http://uploads.invalid",
        region="test-1",
        bucket="test-private",
        access_key_id="fictional-access",
        secret_access_key="fictional-secret",  # noqa: S106 - isolated test credential
    )


@pytest.mark.asyncio
async def test_s3_presigns_bound_put_and_safe_owner_review_get() -> None:
    signer = _Signer()
    storage = object.__new__(AttachmentS3ObjectStorage)
    storage._options = _options()
    storage._signer = signer
    expires_at = datetime.now(UTC) + timedelta(minutes=5)

    upload = await storage.presign_put(
        "opaque/staging/key", AttachmentMediaType.PDF.value, 12_345, expires_at
    )
    download = await storage.presign_get(
        "opaque/quarantine/key",
        'proof "one".pdf',
        AttachmentDownloadPurpose.OWNER_EVIDENCE_REVIEW,
        expires_at,
    )

    assert dict(upload.required_headers) == {
        "content-type": AttachmentMediaType.PDF.value,
        "content-length": "12345",
        "x-amz-meta-expected-size": "12345",
    }
    assert signer.calls[0][1]["Params"]["ContentLength"] == 12_345
    get_params = signer.calls[1][1]["Params"]
    assert get_params["ResponseCacheControl"] == "private, no-store"
    assert "%22one%22" in get_params["ResponseContentDisposition"]
    assert download.required_headers == ()


@pytest.mark.asyncio
async def test_s3_bounded_reads_close_streams_and_return_digest(tmp_path: Path) -> None:
    bodies = [_ClosingBody(b"prefix"), _ClosingBody(b"download")]
    storage = object.__new__(AttachmentS3ObjectStorage)
    storage._options = _options()
    storage._client = _ObjectClient(bodies)

    assert await storage.read_prefix("opaque", 3) == b"pre"
    destination = tmp_path / "attachment.bin"
    downloaded = await storage.download("opaque", destination, 64)

    assert downloaded.path == destination.resolve()
    assert downloaded.size_bytes == len(b"download")
    assert len(downloaded.sha256_digest) == 32
    assert all(body.closed for body in bodies)


@pytest.mark.asyncio
async def test_clamav_transport_outage_returns_unavailable_without_raw_error(
    tmp_path: Path,
) -> None:
    class UnavailableScanner(AttachmentClamAvScanner):
        async def _scan(self, path: Path) -> object:
            del path
            raise OSError("fictional scanner details")

    path = tmp_path / "attachment.bin"
    path.write_bytes(b"safe fixture")
    result = await UnavailableScanner(AttachmentClamAvOptions()).scan(path)

    assert result.verdict is ScanVerdict.UNAVAILABLE


@pytest.mark.asyncio
async def test_bounded_extractor_returns_counts_only_and_rejects_wrong_bytes(
    tmp_path: Path,
) -> None:
    path = tmp_path / "evidence.docx"
    document = Document()
    document.add_heading("Fictional Evidence", level=1)
    document.add_paragraph("A deliberately fictional supporting statement.")
    document.save(path)
    extractor = BoundedAttachmentExtractor()

    result = await extractor.extract(
        path,
        AttachmentMediaType.DOCX,
        AttachmentLimits(temp_root=tmp_path.resolve()),
    )

    assert result.format_valid
    assert result.extracted_characters > 0
    assert result.extracted_blocks == 2
    assert not hasattr(result, "plain_text")

    invalid = tmp_path / "invalid.pdf"
    invalid.write_bytes(b"not a PDF")
    with pytest.raises(UnsafeAttachment) as raised:
        await extractor.extract(
            invalid,
            AttachmentMediaType.PDF,
            AttachmentLimits(temp_root=tmp_path.resolve()),
        )
    assert raised.value.code is SafeAttachmentError.INVALID_DOCUMENT_STRUCTURE
