"""Bounded local PDF/DOCX validation and deterministic text extraction."""

import asyncio
import re
import stat
from collections.abc import Iterable
from pathlib import Path, PurePosixPath
from uuid import UUID, uuid5
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from careeros.modules.resume_health.application.models import (
    DocumentLimits,
    ExtractedBlock,
    ExtractionResult,
    SourceSpanView,
)
from careeros.modules.resume_health.domain import ResumeMediaType
from careeros.modules.resume_health.domain.errors import UnsafeDocument
from careeros.modules.resume_health.infrastructure.layout import analyze_local_layout

PARSER_VERSION = "careeros-local-parser/1.0.0"
_BLOCK_NAMESPACE = UUID("5f80ce6a-096a-44e9-b4d1-335f2da30c32")
_BIDI_CONTROLS = dict.fromkeys(
    [
        *range(0x202A, 0x202F),
        *range(0x2066, 0x206A),
        0x200E,
        0x200F,
        0x061C,
    ]
)
_COLUMN_GAP = re.compile(r"\S[ \t]{8,}\S")
_BULLET_PREFIX = re.compile(r"^(?:[-*•▪◦]|\d+[.)])\s+")


class LocalDocumentExtractor:
    async def extract(
        self, path: Path, media_type: str, limits: DocumentLimits
    ) -> ExtractionResult:
        try:
            return await asyncio.wait_for(
                asyncio.to_thread(self._extract, path, media_type, limits),
                limits.processing_timeout_seconds,
            )
        except TimeoutError as exc:
            raise UnsafeDocument("document_processing_timeout") from exc

    def _extract(self, path: Path, media_type: str, limits: DocumentLimits) -> ExtractionResult:
        return extract_local_document(path, media_type, limits)


class LocalDocumentTextExtractor:
    async def extract_text(
        self, path: Path, media_type: str, limits: DocumentLimits
    ) -> ExtractionResult:
        try:
            return await asyncio.wait_for(
                asyncio.to_thread(extract_local_text, path, media_type, limits),
                limits.processing_timeout_seconds,
            )
        except TimeoutError as exc:
            raise UnsafeDocument("document_processing_timeout") from exc


def extract_local_document(path: Path, media_type: str, limits: DocumentLimits) -> ExtractionResult:
    """Run local text and layout adapters inside the caller's isolation boundary."""
    return analyze_local_layout(extract_local_text(path, media_type, limits))


def extract_local_text(path: Path, media_type: str, limits: DocumentLimits) -> ExtractionResult:
    """Run deterministic local text extraction without layout policy."""
    if path.stat().st_size > limits.max_upload_bytes:
        raise UnsafeDocument("upload_too_large")
    if media_type == ResumeMediaType.PDF.value:
        return _extract_pdf(path, limits)
    if media_type == ResumeMediaType.DOCX.value:
        return _extract_docx(path, limits)
    raise UnsafeDocument("unsupported_document_type")


def _extract_pdf(path: Path, limits: DocumentLimits) -> ExtractionResult:
    raw = path.read_bytes()
    raw_head = raw[:1024]
    if not raw_head.startswith(b"%PDF-"):
        raise UnsafeDocument("document_signature_mismatch")
    if b"PK\x03\x04" in raw:
        raise UnsafeDocument("polyglot_document_rejected")
    with path.open("rb") as source:
        source.seek(max(0, path.stat().st_size - 2048))
        if b"%%EOF" not in source.read():
            raise UnsafeDocument("malformed_pdf")
    try:
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            raise UnsafeDocument("encrypted_document")
        if not 1 <= len(reader.pages) <= limits.max_pdf_pages:
            raise UnsafeDocument("pdf_page_limit_exceeded")
        page_texts: list[str] = []
        layout_signals: list[str] = []
        for page in reader.pages:
            extracted = (page.extract_text() or "").replace("\x00", "")
            sanitized, bidi_removed = _sanitize_extracted_text(extracted)
            page_texts.append(sanitized)
            if bidi_removed:
                layout_signals.append("bidirectional_controls_present")
            try:
                layout_text = page.extract_text(extraction_mode="layout") or ""
            except (TypeError, ValueError):
                layout_text = ""
            if _COLUMN_GAP.search(layout_text):
                layout_signals.append("multi_column_candidate")
    except UnsafeDocument:
        raise
    except Exception as exc:
        raise UnsafeDocument("malformed_pdf") from exc
    plain_text = "\n\f\n".join(page_texts).strip()
    if len(plain_text) > limits.max_extracted_characters:
        raise UnsafeDocument("extracted_text_limit_exceeded")
    blocks = _blocks_from_pages(page_texts, limits.max_extracted_blocks)
    image_only = len(plain_text.strip()) < 40
    warnings = ("image_only_pdf",) if image_only else ()
    return ExtractionResult(
        plain_text=plain_text,
        reading_order=blocks,
        page_count=len(page_texts),
        image_only=image_only,
        warnings=warnings,
        parser_version=PARSER_VERSION,
        layout_signals=tuple(dict.fromkeys(layout_signals)),
    )


def _extract_docx(path: Path, limits: DocumentLimits) -> ExtractionResult:
    raw = path.read_bytes()
    if not raw.startswith(b"PK\x03\x04"):
        raise UnsafeDocument("document_signature_mismatch")
    if b"%PDF-" in raw:
        raise UnsafeDocument("polyglot_document_rejected")
    try:
        with ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_archive_entries:
                raise UnsafeDocument("archive_entry_limit_exceeded")
            total_uncompressed = 0
            total_compressed = 0
            names: set[str] = set()
            for info in infos:
                normalized = PurePosixPath(info.filename.replace("\\", "/"))
                if normalized.is_absolute() or ".." in normalized.parts:
                    raise UnsafeDocument("archive_path_traversal")
                mode = info.external_attr >> 16
                if mode and stat.S_ISLNK(mode):
                    raise UnsafeDocument("archive_symlink_rejected")
                if info.flag_bits & 0x1:
                    raise UnsafeDocument("encrypted_document")
                total_uncompressed += info.file_size
                total_compressed += max(1, info.compress_size)
                names.add(info.filename.casefold())
            if total_uncompressed > limits.max_archive_uncompressed_bytes:
                raise UnsafeDocument("archive_expansion_limit_exceeded")
            if total_uncompressed > total_compressed * limits.max_archive_ratio:
                raise UnsafeDocument("archive_ratio_limit_exceeded")
            if "word/vbaproject.bin" in names or any("macroenabled" in name for name in names):
                raise UnsafeDocument("macro_document_rejected")
            if "[content_types].xml" not in names or "word/document.xml" not in names:
                raise UnsafeDocument("invalid_docx_package")
            if archive.testzip() is not None:
                raise UnsafeDocument("malformed_docx")
    except UnsafeDocument:
        raise
    except (BadZipFile, OSError) as exc:
        raise UnsafeDocument("malformed_docx") from exc

    try:
        document = Document(str(path))
        layout_signals: list[str] = []
        values: list[tuple[str, str]] = []
        for paragraph in document.paragraphs:
            sanitized, bidi_removed = _sanitize_extracted_text(paragraph.text.replace("\x00", ""))
            values.append(("text", sanitized.strip()))
            if bidi_removed:
                layout_signals.append("bidirectional_controls_present")
        for table in document.tables:
            for row in table.rows:
                raw_value = " | ".join(cell.text.replace("\x00", "").strip() for cell in row.cells)
                value, bidi_removed = _sanitize_extracted_text(raw_value)
                if bidi_removed:
                    layout_signals.append("bidirectional_controls_present")
                if value.strip(" |"):
                    values.append(("table", value))
        if any(
            paragraph.text.strip()
            for section in document.sections
            for part in (section.header, section.footer)
            for paragraph in part.paragraphs
        ):
            layout_signals.append("header_footer_present")
    except Exception as exc:
        raise UnsafeDocument("malformed_docx") from exc
    values = [(kind, value) for kind, value in values if value]
    paragraphs = [value for _, value in values]
    plain_text = "\n".join(paragraphs)
    if len(plain_text) > limits.max_extracted_characters:
        raise UnsafeDocument("extracted_text_limit_exceeded")
    blocks = _blocks_from_docx_values(values, limits.max_extracted_blocks)
    return ExtractionResult(
        plain_text=plain_text,
        reading_order=blocks,
        page_count=1,
        image_only=len(plain_text.strip()) < 40,
        warnings=("limited_extractable_text",) if len(plain_text.strip()) < 40 else (),
        parser_version=PARSER_VERSION,
        layout_signals=tuple(dict.fromkeys(layout_signals)),
    )


def _blocks_from_pages(page_texts: Iterable[str], max_blocks: int) -> tuple[ExtractedBlock, ...]:
    blocks: list[ExtractedBlock] = []
    global_offset = 0
    for page_number, page_text in enumerate(page_texts, start=1):
        cursor = 0
        for raw_line in page_text.splitlines():
            text = raw_line.strip()
            if not text:
                cursor += len(raw_line) + 1
                continue
            relative = page_text.find(text, cursor)
            relative = cursor if relative < 0 else relative
            cursor = relative + len(text)
            kind = "bullet" if _BULLET_PREFIX.match(text) else _guess_block_kind(text)
            if len(blocks) >= max_blocks:
                raise UnsafeDocument("extracted_block_limit_exceeded")
            blocks.append(
                ExtractedBlock(
                    kind=kind,
                    text=text,
                    confidence_basis_points=9_000,
                    spans=(
                        SourceSpanView(
                            page=page_number,
                            start=global_offset + relative,
                            end=global_offset + relative + len(text),
                        ),
                    ),
                )
            )
        global_offset += len(page_text) + 3
    return tuple(blocks)


def _blocks_from_docx_values(
    values: list[tuple[str, str]],
    max_blocks: int,
) -> tuple[ExtractedBlock, ...]:
    blocks: list[ExtractedBlock] = []
    offset = 0
    for source_kind, value in values:
        if len(blocks) >= max_blocks:
            raise UnsafeDocument("extracted_block_limit_exceeded")
        kind = (
            "table"
            if source_kind == "table"
            else ("bullet" if _BULLET_PREFIX.match(value) else _guess_block_kind(value))
        )
        blocks.append(
            ExtractedBlock(
                kind=kind,
                text=value,
                confidence_basis_points=8_500 if kind == "table" else 9_000,
                spans=(SourceSpanView(page=1, start=offset, end=offset + len(value)),),
            )
        )
        offset += len(value) + 1
    return tuple(blocks)


def _guess_block_kind(text: str) -> str:
    if len(text) <= 64 and (text.isupper() or text.endswith(":")):
        return "heading"
    return "paragraph"


def stable_block_id(document_id: UUID, block: ExtractedBlock, index: int) -> UUID:
    span = block.spans[0] if block.spans else None
    material = f"{document_id}:{index}:{span.page if span else 0}:{span.start if span else 0}"
    return uuid5(_BLOCK_NAMESPACE, material)


def _sanitize_extracted_text(value: str) -> tuple[str, bool]:
    sanitized = value.translate(_BIDI_CONTROLS)
    return sanitized, sanitized != value
