"""Deterministic, constrained, cross-format resume renderers."""

from __future__ import annotations

import html
import json
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

import reportlab
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
)

from careeros.modules.resume_builder.application import RenderedResume
from careeros.modules.resume_builder.application.ports import ResumeRenderer
from careeros.modules.resume_builder.domain import (
    ResumeEntityFact,
    ResumeFidelityManifest,
    ResumeFontFamily,
    ResumeFormat,
    ResumeLineSpacing,
    ResumeMarginSize,
    ResumeSection,
    ResumeTemplate,
    ResumeVersion,
    build_fidelity_manifest,
    entity_display_lines,
    fidelity_manifest_payload,
    ordered_personal_facts,
    ordered_sections,
)

RENDERER_VERSION = "resume-renderer-v2"
_SANS_REGULAR = "CareerOSVera"
_SANS_BOLD = "CareerOSVeraBold"
_SANS_ITALIC = "CareerOSVeraItalic"


@dataclass(frozen=True, slots=True)
class _TemplateStyle:
    accent: str
    name_alignment: str
    name_size_delta: int
    section_case: str
    section_rule: str
    section_space: float
    bullet_indent: float
    entity_emphasis: str


_TEMPLATE_STYLES = {
    ResumeTemplate.STANDARD_PROFESSIONAL: _TemplateStyle(
        accent="#234E70",
        name_alignment="left",
        name_size_delta=5,
        section_case="title",
        section_rule="below",
        section_space=7,
        bullet_indent=14,
        entity_emphasis="title",
    ),
    ResumeTemplate.COMPACT_TECHNICAL: _TemplateStyle(
        accent="#0F766E",
        name_alignment="left",
        name_size_delta=3,
        section_case="upper",
        section_rule="none",
        section_space=4,
        bullet_indent=10,
        entity_emphasis="organization",
    ),
    ResumeTemplate.EXECUTIVE: _TemplateStyle(
        accent="#1E293B",
        name_alignment="center",
        name_size_delta=8,
        section_case="upper",
        section_rule="above_below",
        section_space=10,
        bullet_indent=16,
        entity_emphasis="title",
    ),
    ResumeTemplate.GRADUATE: _TemplateStyle(
        accent="#4338CA",
        name_alignment="center",
        name_size_delta=5,
        section_case="title",
        section_rule="below",
        section_space=8,
        bullet_indent=14,
        entity_emphasis="organization",
    ),
    ResumeTemplate.CONSULTING_FINANCE: _TemplateStyle(
        accent="#111827",
        name_alignment="left",
        name_size_delta=4,
        section_case="upper",
        section_rule="below",
        section_space=5,
        bullet_indent=12,
        entity_emphasis="title",
    ),
}


class DeterministicResumeRenderer(ResumeRenderer):
    """Render five distinct, single-column, ATS-oriented template families."""

    def render(self, version: ResumeVersion, *, fmt: str) -> RenderedResume:
        requested = ResumeFormat(fmt)
        manifest = build_fidelity_manifest(version)
        expected = tuple(entry.text for entry in manifest.entries)
        reading_order = expected
        if requested is ResumeFormat.PDF:
            content = _pdf(version)
            page_count = len(PdfReader(BytesIO(content), strict=True).pages)
            return RenderedResume(
                media_type="application/pdf",
                filename="resume.pdf",
                content=content,
                expected_lines=expected,
                renderer_version=RENDERER_VERSION,
                reading_order=reading_order,
                page_count=page_count,
            )
        if requested is ResumeFormat.DOCX:
            content = _docx(version)
            return RenderedResume(
                media_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
                filename="resume.docx",
                content=content,
                expected_lines=expected,
                renderer_version=RENDERER_VERSION,
                reading_order=reading_order,
                page_count=_estimated_page_count(version),
            )
        if requested is ResumeFormat.JSON:
            content = _json(version, manifest)
            return RenderedResume(
                media_type="application/json",
                filename="resume.json",
                content=content,
                expected_lines=expected,
                renderer_version=RENDERER_VERSION,
                reading_order=reading_order,
                page_count=_estimated_page_count(version),
            )
        lines = _presentation_lines(version)
        return RenderedResume(
            media_type="text/plain; charset=utf-8",
            filename="resume.txt",
            content=("\n".join(lines) + "\n").encode("utf-8"),
            expected_lines=expected,
            renderer_version=RENDERER_VERSION,
            reading_order=reading_order,
            page_count=_estimated_page_count(version),
        )


def _presentation_lines(version: ResumeVersion) -> tuple[str, ...]:
    lines: list[str] = []
    personal = ordered_personal_facts(version)
    names = [fact.value for fact in personal if fact.kind == "name"]
    contacts = [fact.value for fact in personal if fact.kind != "name"]
    lines.extend(names)
    if contacts:
        lines.append(" | ".join(contacts))
    if version.target_role:
        lines.append(version.target_role)
    entities = {entity.id: entity for entity in version.entities}
    emitted: set[Any] = set()
    for section in ordered_sections(version):
        lines.append(section.title)
        for item in section.items:
            entity = entities.get(item.entity_id) if item.entity_id is not None else None
            if entity is not None and entity.id not in emitted:
                lines.extend(entity_display_lines(entity))
                emitted.add(entity.id)
            lines.append(item.text)
    return tuple(value for value in lines if value.strip())


def _pdf(version: ResumeVersion) -> bytes:
    _register_pdf_fonts()
    layout = version.layout
    style = _TEMPLATE_STYLES[version.template]
    page_size = LETTER if layout.page_size.value == "letter" else A4
    margin = _margin_points(layout.margins)
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title="Resume",
        author="CareerOS",
        subject="Verified resume export",
        creator=RENDERER_VERSION,
        pageCompression=1,
    )
    story = _pdf_story(version, style)
    document.build(story)
    return output.getvalue()


def _pdf_story(version: ResumeVersion, template: _TemplateStyle) -> list[Any]:
    regular, bold, italic = _pdf_fonts(version.layout.font_family)
    base_size = version.layout.font_size_pt
    leading = _leading(base_size, version.layout.line_spacing)
    alignment = TA_CENTER if template.name_alignment == "center" else TA_LEFT
    body = ParagraphStyle(
        "CareerOSBody",
        fontName=regular,
        fontSize=base_size,
        leading=leading,
        textColor=colors.HexColor("#111827"),
        spaceAfter=max(1, leading * 0.18),
        allowWidows=0,
        allowOrphans=0,
    )
    name = ParagraphStyle(
        "CareerOSName",
        parent=body,
        fontName=bold,
        fontSize=base_size + template.name_size_delta,
        leading=leading + template.name_size_delta,
        alignment=alignment,
        textColor=colors.HexColor(template.accent),
        spaceAfter=3,
    )
    contact = ParagraphStyle(
        "CareerOSContact",
        parent=body,
        fontSize=max(8, base_size - 1),
        leading=max(10, leading - 1),
        alignment=alignment,
        textColor=colors.HexColor("#374151"),
        spaceAfter=2,
    )
    target = ParagraphStyle(
        "CareerOSTarget",
        parent=body,
        fontName=italic,
        alignment=alignment,
        textColor=colors.HexColor("#374151"),
        spaceAfter=template.section_space,
    )
    section_heading = ParagraphStyle(
        "CareerOSSection",
        parent=body,
        fontName=bold,
        fontSize=base_size + (2 if version.template is ResumeTemplate.EXECUTIVE else 1),
        leading=leading + 1,
        textColor=colors.HexColor(template.accent),
        spaceBefore=template.section_space,
        spaceAfter=3,
    )
    entity = ParagraphStyle(
        "CareerOSEntity",
        parent=body,
        fontName=bold,
        spaceBefore=3,
        spaceAfter=1,
    )
    detail = ParagraphStyle(
        "CareerOSDetail",
        parent=body,
        fontName=italic,
        fontSize=max(8, base_size - 1),
        leading=max(10, leading - 1),
        textColor=colors.HexColor("#374151"),
        spaceAfter=1,
    )
    bullet = ParagraphStyle(
        "CareerOSBullet",
        parent=body,
        leftIndent=template.bullet_indent,
        firstLineIndent=-8,
        bulletIndent=template.bullet_indent - 8,
        spaceAfter=2 if version.template is not ResumeTemplate.COMPACT_TECHNICAL else 1,
    )
    story: list[Any] = []
    personal = ordered_personal_facts(version)
    for fact in personal:
        if fact.kind == "name":
            story.append(Paragraph(_escaped(fact.value), name))
    contacts = [fact.value for fact in personal if fact.kind != "name"]
    if contacts:
        story.append(Paragraph(_escaped(" | ".join(contacts)), contact))
    if version.target_role:
        story.append(Paragraph(_escaped(version.target_role), target))
    entities = {item.id: item for item in version.entities}
    emitted: set[Any] = set()
    for section in ordered_sections(version):
        heading = _section_heading(section.title, template.section_case)
        if template.section_rule in {"above_below"}:
            story.append(_rule(template.accent, 0.8))
        story.append(Paragraph(_escaped(heading), section_heading))
        if template.section_rule in {"below", "above_below"}:
            story.append(_rule(template.accent, 0.45))
        for item in section.items:
            entity_fact = entities.get(item.entity_id) if item.entity_id is not None else None
            if entity_fact is not None and entity_fact.id not in emitted:
                story.extend(
                    _pdf_entity(
                        entity_fact,
                        entity,
                        detail,
                        emphasize_organization=template.entity_emphasis == "organization",
                    )
                )
                emitted.add(entity_fact.id)
            story.append(Paragraph(_escaped(item.text), bullet, bulletText="•"))
    return story


def _pdf_entity(
    value: ResumeEntityFact,
    entity_style: ParagraphStyle,
    detail_style: ParagraphStyle,
    *,
    emphasize_organization: bool,
) -> list[Any]:
    lines = entity_display_lines(value)
    if not lines:
        return []
    title = lines[0]
    organization = value.organization
    primary = organization if emphasize_organization and organization else title
    details = [line for line in lines if line != primary]
    return [
        KeepTogether(
            [
                Paragraph(_escaped(primary), entity_style),
                *[Paragraph(_escaped(line), detail_style) for line in details],
            ]
        )
    ]


def _docx(version: ResumeVersion) -> bytes:
    document = Document()
    section = document.sections[0]
    _configure_docx_page(section, version)
    style = _TEMPLATE_STYLES[version.template]
    _configure_docx_styles(document, version, style)
    personal = ordered_personal_facts(version)
    alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if style.name_alignment == "center" else WD_ALIGN_PARAGRAPH.LEFT
    )
    for fact in personal:
        if fact.kind == "name":
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = alignment
            paragraph.add_run(fact.value)
    contacts = [fact.value for fact in personal if fact.kind != "name"]
    if contacts:
        paragraph = document.add_paragraph(" | ".join(contacts), style="Subtitle")
        paragraph.alignment = alignment
    if version.target_role:
        paragraph = document.add_paragraph(version.target_role, style="Subtitle")
        paragraph.alignment = alignment
    entities = {item.id: item for item in version.entities}
    emitted: set[Any] = set()
    for resume_section in ordered_sections(version):
        heading = document.add_paragraph(
            _section_heading(resume_section.title, style.section_case),
            style="Heading 1",
        )
        if style.section_rule != "none":
            _docx_paragraph_rule(heading, style.accent)
        _add_docx_section(document, resume_section, entities, emitted, style)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _configure_docx_page(section: Any, version: ResumeVersion) -> None:
    if version.layout.page_size.value == "a4":
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    margin = _margin_inches(version.layout.margins)
    section.top_margin = Inches(margin)
    section.bottom_margin = Inches(margin)
    section.left_margin = Inches(margin)
    section.right_margin = Inches(margin)
    section.start_type = WD_SECTION.NEW_PAGE


def _configure_docx_styles(
    document: Any,
    version: ResumeVersion,
    template: _TemplateStyle,
) -> None:
    family = _docx_font(version)
    normal = document.styles["Normal"]
    normal.font.name = family
    normal.font.size = Pt(version.layout.font_size_pt)
    normal.paragraph_format.line_spacing = _docx_line_spacing(version.layout.line_spacing)
    normal.paragraph_format.space_after = Pt(1)
    title = document.styles["Title"]
    title.font.name = family
    title.font.bold = True
    title.font.size = Pt(version.layout.font_size_pt + template.name_size_delta)
    title.font.color.rgb = _rgb(template.accent)
    subtitle = document.styles["Subtitle"]
    subtitle.font.name = family
    subtitle.font.size = Pt(max(8, version.layout.font_size_pt - 1))
    subtitle.font.italic = True
    heading = document.styles["Heading 1"]
    heading.font.name = family
    heading.font.bold = True
    heading.font.size = Pt(version.layout.font_size_pt + 1)
    heading.font.color.rgb = _rgb(template.accent)
    heading.paragraph_format.space_before = Pt(template.section_space)
    heading.paragraph_format.space_after = Pt(2)


def _add_docx_section(
    document: Any,
    section: ResumeSection,
    entities: dict[Any, ResumeEntityFact],
    emitted: set[Any],
    template: _TemplateStyle,
) -> None:
    for item in section.items:
        entity = entities.get(item.entity_id) if item.entity_id is not None else None
        if entity is not None and entity.id not in emitted:
            lines = entity_display_lines(entity)
            if lines:
                primary_index = (
                    1 if template.entity_emphasis == "organization" and len(lines) > 1 else 0
                )
                for index, line in enumerate(lines):
                    paragraph = document.add_paragraph()
                    run = paragraph.add_run(line)
                    run.bold = index == primary_index
                    run.italic = index != primary_index
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.paragraph_format.space_after = Pt(0)
            emitted.add(entity.id)
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item.text)
        paragraph.paragraph_format.left_indent = Pt(template.bullet_indent)
        paragraph.paragraph_format.space_after = Pt(
            1 if template is _TEMPLATE_STYLES[ResumeTemplate.COMPACT_TECHNICAL] else 2
        )


def _json(version: ResumeVersion, manifest: ResumeFidelityManifest) -> bytes:
    content = {
        "schemaVersion": "resume-export-json-v3",
        "versionId": str(version.id),
        "resumeId": str(version.resume_id),
        "versionNumber": version.version_number,
        "title": version.title,
        "targetRole": version.target_role,
        "template": version.template.value,
        "layout": {
            "fontFamily": version.layout.font_family.value,
            "fontSizePt": version.layout.font_size_pt,
            "lineSpacing": version.layout.line_spacing.value,
            "margins": version.layout.margins.value,
            "pageLimit": version.layout.page_limit,
            "pageSize": version.layout.page_size.value,
        },
        "personalFacts": [
            {
                "id": str(fact.id),
                "kind": fact.kind,
                "value": fact.value,
                "label": fact.label,
                "isPrimary": fact.is_primary,
            }
            for fact in version.personal_facts
        ],
        "entities": [
            {
                "id": str(entity.id),
                "kind": entity.kind,
                "title": entity.title,
                "organization": entity.organization,
                "officialTitle": entity.official_title,
                "displayTitle": entity.display_title,
                "location": entity.location,
                "startDate": (
                    {"year": entity.start_date.year, "month": entity.start_date.month}
                    if entity.start_date is not None
                    else None
                ),
                "endDate": (
                    {"year": entity.end_date.year, "month": entity.end_date.month}
                    if entity.end_date is not None
                    else None
                ),
                "isCurrent": entity.is_current,
                "evidenceIds": [str(value) for value in entity.evidence_ids],
            }
            for entity in version.entities
        ],
        "sections": [
            {
                "id": str(section.id),
                "title": section.title,
                "kind": section.kind,
                "items": [
                    {
                        "id": str(item.id),
                        "text": item.text,
                        "evidenceIds": [str(value) for value in item.evidence_ids],
                        "entityId": str(item.entity_id) if item.entity_id is not None else None,
                        "source": item.source,
                        "evidenceReferences": [
                            {
                                "evidenceId": str(reference.evidence_id),
                                "evidenceRevisionId": str(reference.evidence_revision_id),
                                "revisionNumber": reference.revision_number,
                                "statementSha256": reference.statement_sha256,
                                "claimSha256": reference.claim_sha256,
                                "linkBasis": reference.link_basis.value,
                                "sourceSkillId": (
                                    str(reference.source_skill_id)
                                    if reference.source_skill_id is not None
                                    else None
                                ),
                            }
                            for reference in item.evidence_references
                        ],
                    }
                    for item in section.items
                ],
            }
            for section in ordered_sections(version)
        ],
        "fidelityManifest": fidelity_manifest_payload(manifest),
    }
    return (json.dumps(content, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode(
        "utf-8"
    )


def _estimated_page_count(version: ResumeVersion) -> int:
    width = 92 if version.layout.page_size.value == "letter" else 96
    width -= 14 if version.layout.margins is ResumeMarginSize.WIDE else 0
    width += 12 if version.layout.margins is ResumeMarginSize.NARROW else 0
    units = 5
    for line in _presentation_lines(version):
        units += max(1, (len(line) + width - 1) // width)
    per_page = {
        ResumeLineSpacing.COMPACT: 55,
        ResumeLineSpacing.STANDARD: 47,
        ResumeLineSpacing.RELAXED: 40,
    }[version.layout.line_spacing]
    per_page -= max(0, version.layout.font_size_pt - 10) * 3
    return max(1, (units + per_page - 1) // per_page)


def _register_pdf_fonts() -> None:
    if _SANS_REGULAR in pdfmetrics.getRegisteredFontNames():
        return
    fonts = Path(reportlab.__file__).resolve().parent / "fonts"
    pdfmetrics.registerFont(TTFont(_SANS_REGULAR, str(fonts / "Vera.ttf")))
    pdfmetrics.registerFont(TTFont(_SANS_BOLD, str(fonts / "VeraBd.ttf")))
    pdfmetrics.registerFont(TTFont(_SANS_ITALIC, str(fonts / "VeraIt.ttf")))


def _pdf_fonts(family: ResumeFontFamily) -> tuple[str, str, str]:
    if family is ResumeFontFamily.SERIF:
        return ("Times-Roman", "Times-Bold", "Times-Italic")
    return (_SANS_REGULAR, _SANS_BOLD, _SANS_ITALIC)


def _docx_font(version: ResumeVersion) -> str:
    if version.template is ResumeTemplate.COMPACT_TECHNICAL:
        return "Aptos"
    if version.layout.font_family is ResumeFontFamily.SERIF:
        return "Georgia" if version.template is ResumeTemplate.GRADUATE else "Times New Roman"
    return "Arial"


def _section_heading(value: str, case: str) -> str:
    if case == "upper":
        return value.upper()
    return value.title() if case == "title" else value


def _leading(size: int, spacing: ResumeLineSpacing) -> float:
    multiplier = {
        ResumeLineSpacing.COMPACT: 1.12,
        ResumeLineSpacing.STANDARD: 1.25,
        ResumeLineSpacing.RELAXED: 1.4,
    }[spacing]
    return size * multiplier


def _docx_line_spacing(spacing: ResumeLineSpacing) -> float:
    return {
        ResumeLineSpacing.COMPACT: 1.0,
        ResumeLineSpacing.STANDARD: 1.08,
        ResumeLineSpacing.RELAXED: 1.2,
    }[spacing]


def _margin_points(size: ResumeMarginSize) -> float:
    return _margin_inches(size) * float(inch)


def _margin_inches(size: ResumeMarginSize) -> float:
    return {
        ResumeMarginSize.NARROW: 0.5,
        ResumeMarginSize.STANDARD: 0.7,
        ResumeMarginSize.WIDE: 0.9,
    }[size]


def _rule(color: str, width: float) -> HRFlowable:
    return HRFlowable(
        width="100%",
        thickness=width,
        color=colors.HexColor(color),
        spaceBefore=0,
        spaceAfter=2,
    )


def _docx_paragraph_rule(paragraph: Any, color: str) -> None:
    # python-docx has no public paragraph-border API. A short text rule remains
    # searchable, single-column, and safely decorative rather than carrying facts.
    run = paragraph.add_run("  " + "―" * 8)
    run.font.color.rgb = _rgb(color)


def _rgb(value: str) -> RGBColor:
    normalized = value.removeprefix("#")
    return RGBColor(
        int(normalized[0:2], 16),
        int(normalized[2:4], 16),
        int(normalized[4:6], 16),
    )


def _escaped(value: str) -> str:
    return html.escape(value, quote=False).replace("\n", "<br/>")
