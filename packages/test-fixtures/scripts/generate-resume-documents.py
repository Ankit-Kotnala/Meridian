"""Generate deterministic fictional PDF/DOCX fixtures without office software.

The generated files contain only synthetic ``example.test`` data. Keeping their
source here makes the corpus reviewable and reproducible across platforms.
"""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "generated"

LINES = (
    "ALEX RIVERA",
    "alex.rivera@example.test | +1 555 010 0200 | Portland, OR",
    "SUMMARY",
    "Product manager focused on accessible workflow software.",
    "EXPERIENCE",
    "Senior Product Manager | Northstar Labs | 2022 - Present",
    "Led a cross-functional team of 8 to simplify onboarding for new users.",
    "Reduced setup time from 20 minutes to 12 minutes in a measured pilot.",
    "Product Manager | Cedar Systems | 2019 - 2022",
    "Launched an account recovery workflow with design and engineering partners.",
    "EDUCATION",
    "BSc Information Systems | Example State University | 2015 - 2019",
    "SKILLS",
    "Product strategy, user research, roadmap planning, data analysis",
)


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_bytes(lines: tuple[str, ...], *, include_text: bool = True) -> bytes:
    commands = ["BT", "/F1 10 Tf", "48 760 Td", "13 TL"]
    if include_text:
        for index, line in enumerate(lines):
            if index:
                commands.append("T*")
            commands.append(f"({_pdf_escape(line)}) Tj")
    commands.append("ET")
    return _pdf_document("\n".join(commands).encode("ascii"))


def two_column_pdf_bytes() -> bytes:
    left = (
        "ALEX RIVERA",
        "EXPERIENCE",
        "Engineer | Fictional Alpha | 2021 - Present",
        "Built accessible fictional workflows.",
    )
    right = (
        "SKILLS",
        "Python, product strategy, accessibility",
        "EDUCATION",
        "Example State University | 2020",
    )
    commands = ["BT", "/F1 10 Tf", "40 760 Td", "14 TL"]
    for index, line in enumerate(left):
        if index:
            commands.append("T*")
        commands.append(f"({_pdf_escape(line)}) Tj")
    commands.extend(["ET", "BT", "/F1 10 Tf", "320 760 Td", "14 TL"])
    for index, line in enumerate(right):
        if index:
            commands.append("T*")
        commands.append(f"({_pdf_escape(line)}) Tj")
    commands.append("ET")
    return _pdf_document("\n".join(commands).encode("ascii"))


def _pdf_document(stream: bytes) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def _save_normalized_docx(document: Document, path: Path) -> None:
    source = io.BytesIO()
    document.save(source)
    source.seek(0)
    with (
        zipfile.ZipFile(source) as original,
        zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as normalized,
    ):
        for name in sorted(original.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 7, 15, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            normalized.writestr(info, original.read(name))


def write_adversarial_layout_docx(path: Path) -> None:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "FICTIONAL REPEATED HEADER"
    document.sections[0].footer.paragraphs[0].text = "FICTIONAL REPEATED FOOTER"
    document.add_heading("EXPERIENCE", level=1)
    unusual = document.add_paragraph("Engineer | Fictional Alpha | 03/2021 - 2024")
    unusual.runs[0].font.name = "Papyrus"
    document.add_paragraph("Advisor | Fictional Beta | 2022-06 - Present")
    document.add_paragraph(
        "Visible text with a hidden \u202ebidirectional control.",
        style="List Bullet",
    )
    for index in range(3):
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = f"Fictional project {index}"
        table.cell(0, 1).text = f"20{22 + index}"
    _save_normalized_docx(document, path)


def write_long_docx(path: Path) -> None:
    document = Document()
    document.add_heading("EXPERIENCE", level=1)
    for index in range(600):
        document.add_paragraph(
            f"Fictional bounded achievement {index}: " + "x" * 120,
            style="List Bullet",
        )
    _save_normalized_docx(document, path)


def _paragraph(text: str, *, heading: bool = False, bullet: bool = False) -> str:
    properties = []
    if heading:
        properties.append('<w:pStyle w:val="Heading1"/>')
    if bullet:
        properties.append('<w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>')
    props = f"<w:pPr>{''.join(properties)}</w:pPr>" if properties else ""
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return f'<w:p>{props}<w:r><w:t xml:space="preserve">{escaped}</w:t></w:r></w:p>'


def write_docx(path: Path) -> None:
    heading_values = {"SUMMARY", "EXPERIENCE", "EDUCATION", "SKILLS"}
    bullet_values = {LINES[6], LINES[7], LINES[9]}
    body = "".join(
        _paragraph(line, heading=line in heading_values, bullet=line in bullet_values)
        for line in LINES
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}<w:sectPr/></w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" '
        'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/docProps/app.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
        "</Types>"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        "</Relationships>"
    )
    app = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
        "<Pages>1</Pages><Application>Rezumi fixture generator</Application></Properties>"
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in (
            ("[Content_Types].xml", content_types),
            ("_rels/.rels", relationships),
            ("word/document.xml", document),
            ("docProps/app.xml", app),
        ):
            info = zipfile.ZipInfo(name, date_time=(2026, 7, 15, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, value)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / "fictional-resume.pdf").write_bytes(pdf_bytes(LINES))
    (OUTPUT / "image-only.pdf").write_bytes(pdf_bytes((), include_text=False))
    (OUTPUT / "two-column.pdf").write_bytes(two_column_pdf_bytes())
    write_docx(OUTPUT / "fictional-resume.docx")
    write_adversarial_layout_docx(OUTPUT / "adversarial-layout.docx")
    write_long_docx(OUTPUT / "long-resume.docx")
    manifest = {
        path.name: {
            "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        }
        for path in sorted(OUTPUT.iterdir())
        if path.is_file() and path.name != "manifest.json"
    }
    (OUTPUT / "manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
